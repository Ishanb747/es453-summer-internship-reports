"""Assemble Report 1 and Report 2 as .docx (forms appended with docxcompose)."""
import sys
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docxcompose.composer import Composer

import doclib as L
import forms
import report1 as R
from data import STUDENT, ORG, WEEKS

OUT = sys.argv[1]
import json, os
PAGES = json.load(open(sys.argv[2])) if len(sys.argv) > 2 and os.path.exists(sys.argv[2]) else {}
SRC = "/Users/ishan/Downloads/Students_Documents_7th_ES453"


def labelled(comp, title, form, in_toc=True):
    """Small heading at the top of the form's own page, then the form."""
    L.form_label(comp.doc, title, in_toc=in_toc)
    comp.append(form)


def logbooks(comp, title):
    for i, w in enumerate(WEEKS):
        if i == 0:
            labelled(comp, title, forms.logbook(w))
        else:
            labelled(comp, f"{title.split('(')[0].strip().title()} – Week {w['no']}", forms.logbook(w), in_toc=False)


def appendices(comp):
    logbooks(comp, "APPENDIX D – WEEKLY LOGBOOKS (WEEKS 1–17)")
    labelled(comp, "APPENDIX F – INDUSTRY MENTOR FEEDBACK FORM", forms.mentor_feedback())
    labelled(comp, "APPENDIX G – STUDENT FEEDBACK FORM", forms.student_feedback())
    labelled(comp, "APPENDIX H – COURSE EXIT SURVEY", forms.exit_survey())
    labelled(comp, "APPENDIX J – CO–PO / PSO MAPPING", forms.co_po())
    labelled(comp, "APPENDIX K – SDG MAPPING", forms.sdg())


def pb(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def report1():
    d = R.build_body("COMPREHENSIVE TRAINING REPORT", PAGES.get("r1"))
    comp = Composer(d)
    appendices(comp)
    L.page_borders(comp.doc)
    comp.save(f"{OUT}/Report1_Comprehensive_Training_Report.docx")


# ------------------------------------------------------------------ Report 2

def anonymise(t):
    for a, b in (("the Sulekal site", "a client site"), ("Sulekal", "a client"), ("Sandoz raw data", "client-site raw data"), ("Sandoz", "client")):
        t = t.replace(a, b)
    return t


def tidy(t):
    parts = [x.strip().rstrip(".") for x in t.split("\n") if x.strip()]
    parts = [x[0].upper() + x[1:] if x else x for x in parts]
    return "; ".join(parts) + "."


def tracker_rows(path):
    ws = openpyxl.load_workbook(path, data_only=True)["IB"]
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        sn, date, day, hrs, task = r[0], r[1], r[2], r[3], r[4]
        if hasattr(date, "strftime") and task:
            rows.append([date.strftime("%d %b %Y"), day, "–" if hrs is None else f"{hrs:g}", tidy(anonymise(str(task)))])
    return rows


from build_specs import REPORT2_INDEX


def report2():
    d = L.new_document()
    R.cover(d, "SUPPORTING DOCUMENTS REPORT")
    L.new_section(d, fmt="lowerRoman", start=1)
    R.index_page(d, REPORT2_INDEX, PAGES.get("r2"), first=True)
    L.new_section(d, fmt="decimal", start=1)
    comp = Composer(d)

    first = L.form_label(d, "1. FINAL SUBMISSION & VIVA DOCUMENT CHECKLIST")
    first.paragraph_format.page_break_before = False
    comp.append(forms.checklist())

    L.h1(d, "2. OFFER / INTERNSHIP AGREEMENT")
    L.para(d, "Internship Agreement between PV Doctor Pte. Ltd. and Ishan Bhattacharya, dated 4 May 2026 (6 pages). The personal identification number on page 2 has been masked.",
           italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=2)
    for i in range(1, 7):
        img = f"{R.PAGES}/agr-2-redacted.png" if i == 2 else f"{R.PAGES}/agr-{i}.png"
        if i > 1:
            pb(d)
        L.full_page_image(d, img, width_cm=14.6 if i == 1 else 15.3)

    L.h1(d, "3. INTERNSHIP COMPLETION CERTIFICATE")
    L.full_page_image(d, f"{R.PAGES}/cert-1.png", width_cm=15.3)

    L.h1(d, "4. NOC STATUS NOTE")
    for t, bold in (("MAHARAJA AGRASEN INSTITUTE OF TECHNOLOGY", True), ("Department of Information Technology", True), ("NOC (No Objection Certificate) – Status Note", True), ("", False),
                    (f"Student Name: {STUDENT['name']}", False), (f"Enrollment No.: {STUDENT['roll']}", False), (f"Counselor Group: {STUDENT['group']}", False),
                    ("Department: Information Technology", False), ("", False),
                    ("Regarding: NOC for Summer Internship/Training at PV Doctor (Singapore, Remote)", True), ("", False),
                    ("This internship was fully remote, and the company (PV Doctor) did not require or ask for a No Objection Certificate as part of onboarding. Hence, no NOC was obtained.", False),
                    ("", False)):
        L.para(d, t, bold=bold, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p = L.para(d, "", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p.add_run().add_picture(R.SIG, height=L.Cm(0.5))
    for t in (STUDENT["name"], STUDENT["roll"]):
        L.para(d, t, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)

    L.h1(d, "5. INDUSTRY MENTOR AND ORGANISATION DETAILS")
    L.table(d, ["Particular", "Details"], [
        ["Organisation", ORG["full"]],
        ["Website / Portal", "www.pv.doctor  |  portal.pv.doctor"],
        ["Department", "Data Engineering & Analytics"],
        ["Industry Mentor", f"Name: {ORG['mentor']}\nDesignation: {ORG['mentor_desig']}\nE-mail / Phone: {ORG['mentor_contact']}"],
        ["HR / Administration contact", "Ms. Dannila Choong, Head of Finance & General Affairs (signatory of the internship certificate)"],
        ["Managing Director", "Dr. André Nobre (signatory of the internship agreement)"],
        ["Company contact (public)", "SmartOM.IT@pv.doctor  |  WhatsApp +65 8896 0180"],
        ["Student role", "Data Analytics Intern – Web Development of the Smart O&M Monitoring Portal"],
        ["Duration / Mode", f"{STUDENT['duration']} ({STUDENT['weeks']} weeks) – Remote"],
        ["Stipend", "INR 20,000 per month"],
    ], widths=[5, 11], title="Mentor and organisation details")

    R.declaration(d, heading="6. DECLARATION")


    L.h1(d, "7. WORK EVIDENCE – INTERN TASK TRACKER EXTRACTS")
    L.para(d, "Daily entries submitted in PV Doctor's intern task tracker. Only the student's own rows are reproduced. Days without an entry (weekends, leave) are omitted.", size=11)
    L.table(d, ["Date", "Day", "Hrs", "Task / activity description"], tracker_rows(f"{SRC}/2606 Intern Task Tracker - June 2026.xlsx"),
            widths=[2.3, 2.0, 1.0, 10.7], title="Task tracker – June 2026", size=9.5)
    L.table(d, ["Date", "Day", "Hrs", "Task / activity description"], tracker_rows(f"{SRC}/2608 Intern Task Tracker - August 2026.xlsx"),
            widths=[2.3, 2.0, 1.0, 10.7], title="Task tracker – August 2026", size=9.5)

    L.h1(d, "8. WORK EVIDENCE – PORTAL SCREENSHOTS")
    L.para(d, "Screenshots from the public demo account of portal.pv.doctor (anonymised demo sites) showing modules developed during the internship.", size=11)
    C2 = f"{R.SP}/cards"
    for path, cap in [(f"{C2}/c_portfolio.jpg", "Portfolio Overview"), (f"{C2}/c_pr_evo.jpg", "Performance Ratio Evolution with corridor"),
                      (f"{C2}/c_irr_yield.jpg", "Irradiation vs Yield"), (f"{C2}/c_inv_heatmap.jpg", "Inverter Performance Heatmap"),
                      (f"{C2}/c_spatial.jpg", "PV Spatial Heatmap (PV Healthscan)")]:
        L.figure(d, path, cap)
    labelled(comp, "ANNEXURE 1 – MID-INTERNSHIP VERIFICATION RECORD", forms.mid_verification())
    logbooks(comp, "ANNEXURE 2 – WEEKLY LOGBOOKS (APPENDIX D)")
    labelled(comp, "ANNEXURE 3 – INDUSTRY MENTOR FEEDBACK FORM (APPENDIX F)", forms.mentor_feedback())
    labelled(comp, "ANNEXURE 4 – STUDENT FEEDBACK FORM (APPENDIX G)", forms.student_feedback())
    labelled(comp, "ANNEXURE 5 – COURSE EXIT SURVEY (APPENDIX H)", forms.exit_survey())
    labelled(comp, "ANNEXURE 6 – CO–PO/PSO MAPPING (APPENDIX J)", forms.co_po())
    labelled(comp, "ANNEXURE 7 – SDG MAPPING (APPENDIX K)", forms.sdg())
    L.page_borders(comp.doc)
    comp.save(f"{OUT}/Report2_Supporting_Documents.docx")


if __name__ == "__main__":
    report1()
    report2()
    print("ok")
