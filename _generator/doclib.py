"""python-docx helpers implementing the ES-453 report format."""
import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def _rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def new_document():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin, s.top_margin, s.right_margin, s.bottom_margin = Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)
    s.footer_distance = Cm(1.3)
    st = d.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    _rfonts(st.element)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(8)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, align, before in (("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0),
                                      ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 18),
                                      ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 12)):
        h = d.styles[name]
        h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
        h.font.italic = False
        h.font.color.rgb = RGBColor(0, 0, 0)
        _rfonts(h.element)
        h.paragraph_format.alignment = align
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(18 if name == "Heading 1" else 8)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            h.paragraph_format.page_break_before = True
    for name in ("List Bullet", "List Number", "Caption"):
        st2 = d.styles[name]
        st2.font.name = FONT
        _rfonts(st2.element)
    cap = d.styles["Caption"]
    cap.font.size, cap.font.bold, cap.font.italic = Pt(11), True, False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(14)
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.line_spacing = 1.15
    return d


# ---------------------------------------------------------------- fields

def _field(paragraph, instr, placeholder=""):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r2._r.append(it)
    r3 = paragraph.add_run()
    sp = OxmlElement("w:fldChar"); sp.set(qn("w:fldCharType"), "separate"); r3._r.append(sp)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r5._r.append(e)
    return r4


def toc(d, instr):
    p = d.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _field(p, instr, "Right-click and choose Update Field to refresh this list.")


# ---------------------------------------------------------------- sections & page numbers

def _pgnum(section, fmt, start=None):
    sp = section._sectPr
    pg = sp.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sp.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = _field(p, "PAGE", "1")
    for run in p.runs:
        run.font.name, run.font.size = FONT, Pt(11)


def new_section(d, fmt=None, start=None, numbered=True):
    s = d.add_section(WD_SECTION.NEW_PAGE)
    if fmt:
        _pgnum(s, fmt, start)
    if numbered:
        footer_page_number(s)
    return s


# ---------------------------------------------------------------- content helpers

def para(d, text, bold=False, italic=False, align=None, size=None, space_after=None, keep=False):
    p = d.add_paragraph()
    _rich(p, text, bold, italic, size)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def _rich(p, text, bold=False, italic=False, size=None):
    """**bold** and <sub>subscript</sub> segments inside text."""
    import re
    parts = text.split("**")
    for i, part in enumerate(parts):
        for seg in re.split(r"(<sub>.*?</sub>)", part):
            if not seg:
                continue
            sub = seg.startswith("<sub>")
            r = p.add_run(seg[5:-6] if sub else seg)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if sub:
                r.font.subscript = True
            if size:
                r.font.size = Pt(size)


def bullets(d, items, style="List Bullet"):
    for it in items:
        p = d.add_paragraph(style=style)
        _rich(p, it)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5


def h1(d, text):
    return d.add_heading(text, level=1)


def h2(d, text):
    return d.add_heading(text, level=2)


def h3(d, text):
    return d.add_heading(text, level=3)


def caption(d, kind, text):
    p = d.add_paragraph(style="Caption")
    r = p.add_run(f"{kind} ")
    _field(p, f"SEQ {kind} \\* ARABIC", "#")
    p.add_run(f": {text}")
    return p


def figure(d, path, text, width_cm=15.5):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(d, "Figure", text)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(d, header, rows, widths=None, title=None, size=10.5):
    if title:
        cp = caption(d, "Table", title)
        cp.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h); r.bold = True; r.font.size = Pt(size)
        _shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _rich(p, str(v), size=size)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if widths:
                c.width = Cm(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 45), ("left", 100), ("bottom", 45), ("right", 100)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
    for row in t.rows:
        cs = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cs)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    sp = d.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.line_spacing = Pt(14)
    sp.add_run(" ").font.size = Pt(6)
    return t


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def full_page_image(d, path, width_cm=16.0, title=None):
    if title:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title); r.bold = True; r.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width_cm))


def no_outline(p):
    """Keep a heading-styled paragraph out of the table of contents."""
    pPr = p._p.get_or_add_pPr()
    o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), "9"); pPr.append(o)


def form_label(d, text, in_toc=True):
    """Small label at the top of a form page (starts a new page)."""
    if in_toc:
        p = d.add_heading(text, level=1)
    else:
        p = d.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        r.font.size = Pt(11); r.bold = True; r.italic = not in_toc
    return p


def page_borders(doc):
    """Single-line page border on every section, like the department sample."""
    for sec in doc.sections:
        sp = sec._sectPr
        old = sp.find(qn("w:pgBorders"))
        if old is not None:
            sp.remove(old)
        pb = OxmlElement("w:pgBorders"); pb.set(qn("w:offsetFrom"), "page")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "24"); e.set(qn("w:color"), "000000")
            pb.append(e)
        # pgBorders must come after pgMar/pgSz in sectPr
        anchor = sp.find(qn("w:pgMar"))
        anchor.addnext(pb)


def index_line(d, num, text, page, bold=True):
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.line_spacing = 1.0
    pf.space_after = Pt(7)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.tab_stops.add_tab_stop(Cm(15.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(f"{num}.  {text}\t{page}")
    r.bold = bold; r.font.size = Pt(11.5)
    return p
