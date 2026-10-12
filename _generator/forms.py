"""Fill the official ES-453 form templates (docx) with Ishan's details."""
import copy
import docx
import docx.text.paragraph
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING

from data import STUDENT, ORG, WEEKS

MENTOR_DATE = "31 August 2026"

TPL = "/Users/ishan/Downloads/Students_Documents_7th_ES453"
TICK = "☑"


def _size_run(runs):
    for r in runs:
        if r.font.size:
            return r
    return runs[0] if runs else None


def set_par(p, text):
    """Replace a paragraph's text, keeping the formatting of its sized run."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    keep = _size_run(runs)
    for r in runs:
        r.text = ""
    keep.text = text


def set_cell(cell, text, bold=None):
    lines = text.split("\n")
    p = cell.paragraphs[0]
    size = None
    r0 = _size_run(p.runs)
    if r0 is not None and r0.font.size:
        size = r0.font.size
    set_par(p, lines[0])
    for extra in lines[1:]:
        np_ = cell.add_paragraph()
        r = np_.add_run(extra)
        r.font.name = "Times New Roman"
        if size:
            r.font.size = size
    if bold is not None:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = bold


SIG = "/private/tmp/claude-501/-Users-ishan-Documents-New-project-Summer-Intern-report/76f2bf2a-0fb3-4818-a4c4-dc9469e7c347/scratchpad/sig2/ishan_signature.png"


def sign(d, label, height_cm=0.75):
    """Replace the underscore blank right after `label` with the student's signature image."""
    import re
    for p in all_pars(d):
        t = p.text
        i = t.find(label)
        if i < 0:
            continue
        m = re.compile(r"_{4,}").search(t, i + len(label))
        if not m:
            continue
        keep = _size_run(p.runs)
        size, bold = (keep.font.size, keep.bold) if keep is not None else (None, None)
        for r in p.runs:
            r._r.getparent().remove(r._r)
        def add(text):
            r = p.add_run(text); r.font.size = size; r.bold = bold; r.font.name = "Times New Roman"
        add(t[:m.start()])
        p.add_run().add_picture(SIG, height=Cm(height_cm))
        add(t[m.end():])
        return True
    raise KeyError(label)


def all_pars(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def replace(d, old, new, count=None):
    n = 0
    for p in all_pars(d):
        if old in p.text:
            set_par(p, p.text.replace(old, new))
            n += 1
            if count and n >= count:
                break
    return n


def startswith(d, prefix, new):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            set_par(p, new)
            return True
    raise KeyError(prefix)


def tick_line(d, label):
    """Turn '☐ label' into '☑ label' wherever it appears."""
    return replace(d, "☐ " + label, TICK + " " + label) or replace(d, "☐  " + label, TICK + "  " + label)


def fit_tables(d, max_twips=int(15.9 / 2.54 * 1440)):
    """Scale fixed-width template tables (designed for Letter, 1.5 cm margins) to the report text width."""
    from docx.oxml.ns import qn
    for t in d.tables:
        tbl = t._tbl
        grid = tbl.tblGrid
        cols = grid.findall(qn("w:gridCol"))
        total = sum(int(c.get(qn("w:w"))) for c in cols)
        if total <= max_twips:
            continue
        k = max_twips / total
        for c in cols:
            c.set(qn("w:w"), str(int(int(c.get(qn("w:w"))) * k)))
        for tcW in tbl.iter(qn("w:tcW")):
            if tcW.get(qn("w:type")) in (None, "dxa"):
                tcW.set(qn("w:w"), str(int(int(tcW.get(qn("w:w"))) * k)))
        tblPr = tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is not None and tblW.get(qn("w:type")) == "dxa":
            tblW.set(qn("w:w"), str(max_twips))
        ind = tblPr.find(qn("w:tblInd"))
        if ind is not None:
            tblPr.remove(ind)


def compact(d, grow=0.0):
    """Even, readable spacing for a form placed inside the report (A4, report margins)."""
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    fit_tables(d)
    body = d.element.body
    prev_is_table = False
    for el in body.iterchildren():
        if el.tag == qn("w:tbl"):
            prev_is_table = True
            continue
        if el.tag != qn("w:p"):
            continue
        p = docx.text.paragraph.Paragraph(el, d)
        pf = p.paragraph_format
        pf.line_spacing = 1.08
        pf.space_after = Pt(1.5)
        runs = [r for r in p.runs if r.text.strip()]
        heading = bool(runs) and all(r.bold for r in runs) and p.alignment != WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(5) if (prev_is_table or heading) else Pt(0)
        prev_is_table = False
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    if p.alignment in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for p in all_pars(d):
        for r in p.runs:
            if not r.font.name:
                r.font.name = "Times New Roman"
            if grow and r.font.size and r.font.size <= Pt(11):
                r.font.size = Pt(r.font.size.pt + grow)


def info_2col(t):
    """Student / Industry info table used in feedback forms (label: ____ cells)."""
    vals = {
        "Name:": STUDENT["name"],
        "Roll No.:": STUDENT["roll"],
        "Batch / Section:": STUDENT["batch"],
        "Semester:": "7th",
        "Organization Name:": ORG["name"],
        "Contact Person (Mentor/HR):": ORG["mentor"],
        "Designation:": ORG["mentor_desig"],
        "Contact No. & Email:": ORG["mentor_contact"],
    }
    for row in t.rows:
        for c in row.cells:
            txt = c.text.strip()
            for k, v in vals.items():
                if txt.startswith(k):
                    set_cell(c, f"{k} {v}")
                    break


def info_4col(t):
    vals = {
        "Name": STUDENT["name"], "Roll No.": STUDENT["roll"], "Batch / Section": STUDENT["batch"],
        "Semester": "7th", "Project / Domain Name": STUDENT["domain"],
        "Organization Name": ORG["name"], "Organization / Institution Name": ORG["name"],
        "Contact Person (Mentor/HR)": ORG["mentor"], "Contact Person (Mentor/HR/Supervisor)": ORG["mentor"],
        "Designation": ORG["mentor_desig"], "Contact No. & Email": ORG["mentor_contact"],
        "Faculty Counsellor Name": STUDENT["counsellor"],
    }
    for row in t.rows:
        cells = row.cells
        for i in (0, 2):
            label = cells[i].text.strip()
            if label in vals and i + 1 < len(cells):
                set_cell(cells[i + 1], vals[label])


def particulars(t, extra=None):
    vals = {
        "Student Name": STUDENT["name"], "Roll Number": STUDENT["roll"],
        "Training Duration & Mode": f"{STUDENT['duration']} ({STUDENT['weeks']} weeks) – Remote / Work from Home",
        "Project Title / Course Name": STUDENT["title"],
        "Organization / Institution": ORG["full"],
        "Faculty Counsellor": STUDENT["counsellor"],
    }
    if extra:
        vals.update(extra)
    for row in t.rows:
        label = row.cells[0].text.strip()
        if label in vals:
            set_cell(row.cells[1], vals[label])
        if label == "Type of Training":
            c = row.cells[1]
            set_par(c.paragraphs[0], c.text.replace("☐ Industrial Internship", TICK + " Industrial Internship"))


def tick_col(t, row_idx, col_idx):
    set_cell(t.rows[row_idx].cells[col_idx], TICK)


# ---------------------------------------------------------------- forms

def logbook(week):
    d = docx.Document(f"{TPL}/02_Logbooks/ES453_Proj_Ind_Weekly_Logbook.docx")
    startswith(d, "Week No.:", f"Week No.: {week['no']}     Dates: From {week['from']} to {week['to']}")
    info_4col(d.tables[0])
    for p in d.paragraphs:
        if p.text.startswith("☐ "):
            set_par(p, TICK + p.text[1:])
    startswith(d, "Major Activities Completed:", "Major Activities Completed: " + week["activities"])
    startswith(d, "Key Learning Outcomes Achieved:", "Key Learning Outcomes Achieved: " + week["learning"])
    startswith(d, "Plan for Next Week:", "Plan for Next Week: " + week["plan"])
    startswith(d, "Student Signature:", f"Student Signature: __________________________   Date: {week['to']}")
    sign(d, "Student Signature:")
    startswith(d, "Industry Mentor / Research Supervisor Remarks:",
               f"Industry Mentor / Research Supervisor Remarks: {week.get('mentor_remark', 'Work completed as planned; good progress this week.')}")
    for p in d.paragraphs:
        if p.text.strip().startswith("Signature: ___") and "Date: ___" in p.text:
            set_par(p, f"Signature: ______________________   Name: {ORG['mentor']}   Date: {week['to']}")
    compact(d, grow=0.0)
    return d


def mentor_feedback():
    d = docx.Document(f"{TPL}/04_Feedback/ES453_Industry_Mentor_Feedback_Form.docx")
    info_2col(d.tables[0])
    t = d.tables[1]  # cols: param, Excellent(1), Good(2), Satisfactory(3), Poor(4)
    for row, col in zip(range(1, 6), (2, 1, 1, 2, 2)):
        tick_col(t, row, col)
    t2 = d.tables[2]
    set_cell(t2.rows[1].cells[1], "React / Next.js, TypeScript, Chart.js, Plotly.js, FastAPI, SQL Server, Git")
    c = t2.rows[2].cells[1]; set_cell(c, c.text.replace("☐ Good", TICK + " Good"))
    c = t2.rows[3].cells[1]; set_cell(c, c.text.replace("☐ Completed Successfully", TICK + " Completed Successfully"))
    set_cell(t2.rows[4].cells[1], "Learns new tools quickly; careful attention to UI detail and data correctness; responds well to review feedback; took ownership of chart and raw-data features.")
    set_cell(t2.rows[5].cells[1], "Break large changes into smaller pull requests from the start and write automated tests alongside new features.")
    for p in d.paragraphs:
        if p.text.strip().startswith("Signature of Industry Mentor"):
            set_par(p, f"Signature of Industry Mentor: ______________________   Name: {ORG['mentor']}   Date: {MENTOR_DATE}")
    compact(d, grow=0.0)
    return d


def student_feedback():
    d = docx.Document(f"{TPL}/04_Feedback/ES453_Student_Feedback_Form_Internship.docx")
    info_2col(d.tables[0])
    t = d.tables[1]
    # cols: S.No, Param, Excellent(2), Good(3), Satisfactory(4), Poor(5)
    for r, col in zip(range(1, 5), (3, 2, 2, 2)):
        tick_col(t, r, col)
    t2 = d.tables[2]
    set_cell(t2.rows[1].cells[1], "1. React / Next.js with TypeScript   2. Chart.js & Plotly.js data visualisation   3. FastAPI & SQL Server")
    set_cell(t2.rows[2].cells[1], f"{TICK} Yes      ☐ No")
    set_cell(t2.rows[3].cells[1], STUDENT["student_comment"])
    replace(d, "Student's Signature: ______________________________     Date: __________________",
            f"Student's Signature: ______________________________     Date: {STUDENT['sign_date']}")
    sign(d, "Student's Signature:")
    compact(d, grow=0.5)
    return d


def exit_survey():
    d = docx.Document(f"{TPL}/05_Exit_Survey/Exit_Survey_ES453.docx")
    sugg = list(STUDENT["exit_suggestions"])
    t = d.tables[0]
    for r, score in zip(range(1, 5), STUDENT["exit_scores"]):
        c = t.rows[r].cells[2]
        set_cell(c, c.text.replace(f"□ {score}", f"{TICK} {score}"))
    for p in d.paragraphs:
        if p.text.startswith("____") :
            set_par(p, sugg.pop(0) if sugg else p.text)
    startswith(d, "Name of the Student:", f"Name of the Student: {STUDENT['name']}    Enrollment No.: {STUDENT['roll']}")
    startswith(d, "Signature:", f"Signature: ____________________________          Date: {STUDENT['sign_date']}")
    sign(d, "Signature:")
    compact(d, grow=0.5)
    return d


def co_po():
    d = docx.Document(f"{TPL}/06_Mapping/CO_PO_PSO_ES453.docx")
    particulars(d.tables[0])
    replace(d, "Student Signature: ______________________   Date: _______________",
            f"Student Signature: ______________________   Date: {STUDENT['sign_date']}")
    sign(d, "Student Signature:")
    compact(d)
    return d


def sdg():
    d = docx.Document(f"{TPL}/06_Mapping/SDG_ES453.docx")
    particulars(d.tables[0], {"Organization / Institution": "PV Doctor Pte. Ltd., Singapore", "Training Duration & Mode": "5 May – 31 Aug 2026 (17 weeks), Remote"})
    t = d.tables[1]
    for row in t.rows[1:]:
        goal = row.cells[1].text.strip()
        for key, just in STUDENT["sdg"].items():
            if goal.startswith(key + ":"):
                set_cell(row.cells[0], TICK)
                set_cell(row.cells[2], just)
    replace(d, "Student Signature: ______________________   Date: ___________",
            f"Student Signature: ______________________   Date: {STUDENT['sign_date']}")
    sign(d, "Student Signature:")
    compact(d, grow=-1.0)
    return d


def mid_verification():
    d = docx.Document(f"{TPL}/03_Mid_Verification/ES453_Mid_Internship_Verification.docx")
    info_4col(d.tables[0])
    replace(d, "☐ Yes (Active)", TICK + " Yes (Active)")
    replace(d, "☐ Remote / Work from Home", TICK + " Remote / Work from Home")
    replace(d, "☐ Project Execution Phase", TICK + " Project Execution Phase")
    for p in d.paragraphs:
        if p.text.strip().startswith("1. ______________________________   2."):
            set_par(p, "1. React / Next.js (TypeScript)   2. Chart.js & Plotly.js visualisation   3. FastAPI + SQL Server")
    compact(d, grow=0.5)
    return d


def checklist():
    d = docx.Document(f"{TPL}/01_Instructions/Final_Submission_Viva_Checklist_ES-453.docx")
    startswith(d, "Student Name:", f"Student Name: {STUDENT['name']}   Roll No.: {STUDENT['roll']}   Group: {STUDENT['group']}")
    startswith(d, "Faculty Counsellor:", f"Faculty Counsellor: {STUDENT['counsellor'] or '______________________________'}   Track: Industrial Internship")
    # tick the Industrial Internship block and the mandatory requirements
    in_block = False
    for p in d.paragraphs:
        t = p.text.strip()
        if t == "Industrial Internship":
            in_block = True
            continue
        if t in ("Research Internship",):
            in_block = False
        if t == "Mandatory Requirements":
            in_block = True
            continue
        if t.startswith("SWAYAM/NPTEL:"):
            in_block = False
        if in_block and t.startswith("☐"):
            set_par(p, p.text.replace("☐", TICK, 1))
    replace(d, "Student Signature: ____________________   Date: __________",
            f"Student Signature: ____________________   Date: {STUDENT['sign_date']}")
    sign(d, "Student Signature:", height_cm=0.45)
    compact(d, grow=-1.5)
    return d
